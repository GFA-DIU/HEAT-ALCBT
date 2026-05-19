import os
from collections import defaultdict

import openpyxl

from pages.models.epd import EPD


EXCEL_PATH = os.path.join(
    os.path.dirname(__file__),
    "..", "..", "..", "docs",
    "ALL EPDs.xlsx",
)

SHEET = "epds"

COL_ID = "id"
COL_NAME = "name"
COL_MAPPED = "Mapped Correct Name to be Filled by Rohit"


def update_epd_names(stdout=None, style=None):
    def log(msg):
        if stdout:
            stdout.write(msg)

    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)

    if SHEET not in wb.sheetnames:
        raise RuntimeError(f"Sheet '{SHEET}' not found in workbook.")

    ws = wb[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        raise RuntimeError("Sheet is empty.")

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]

    try:
        col_id = headers.index(COL_ID)
        col_name = headers.index(COL_NAME)
        col_mapped = headers.index(COL_MAPPED)
    except ValueError as e:
        raise RuntimeError(f"Missing expected column: {e}")

    # First pass: find UUIDs that appear more than once among rows with a mapped name
    uuid_mapped_names = defaultdict(list)
    for row in rows[1:]:
        mapped_val = row[col_mapped]
        if mapped_val is None or str(mapped_val).strip() == "":
            continue
        id_val = row[col_id]
        if id_val is None or not str(id_val).strip():
            continue
        uuid_str = str(id_val).strip()
        uuid_mapped_names[uuid_str].append(str(mapped_val).strip())

    duplicate_uuids = {uuid for uuid, names in uuid_mapped_names.items() if len(names) > 1}

    if duplicate_uuids:
        log("\n[DUPLICATES] The following UUIDs appear more than once — skipping all of them:")
        for uuid in duplicate_uuids:
            mapped_list = uuid_mapped_names[uuid]
            log(f"  UUID={uuid} has {len(mapped_list)} mapped names: {mapped_list}")
        log("")

    updated = 0
    skipped_no_mapped = 0
    skipped_not_found = 0
    skipped_duplicate = 0

    for row in rows[1:]:
        mapped_val = row[col_mapped]
        if mapped_val is None or str(mapped_val).strip() == "":
            skipped_no_mapped += 1
            continue

        new_name = str(mapped_val).strip()
        id_val = row[col_id]
        name_val = row[col_name]

        # Try by UUID first
        if id_val is not None and str(id_val).strip():
            uuid_str = str(id_val).strip()

            if uuid_str in duplicate_uuids:
                skipped_duplicate += 1
                continue

            try:
                epd = EPD.objects.get(UUID=uuid_str)
                old_name = epd.name
                epd.name = new_name
                epd.save(update_fields=["name"])
                log(f"  [UUID] '{old_name}' -> '{new_name}' (UUID={uuid_str})")
                updated += 1
                continue
            except EPD.DoesNotExist:
                log(f"  [WARN] UUID not found in DB: {uuid_str} — trying name fallback")

        # Fallback: exact match on name column
        if name_val is not None and str(name_val).strip():
            exact_name = str(name_val).strip()
            matches = EPD.objects.filter(name=exact_name)
            count = matches.count()
            if count == 1:
                epd = matches.first()
                epd.name = new_name
                epd.save(update_fields=["name"])
                log(f"  [NAME] '{exact_name}' -> '{new_name}'")
                updated += 1
            elif count > 1:
                log(f"  [SKIP] Multiple EPDs match name='{exact_name}', ambiguous — skipping")
                skipped_not_found += 1
            else:
                log(f"  [SKIP] No EPD found with name='{exact_name}'")
                skipped_not_found += 1
        else:
            skipped_not_found += 1

    log(
        f"\nSummary: updated={updated}, "
        f"skipped_duplicate_uuid={skipped_duplicate}, "
        f"not_found={skipped_not_found}, "
        f"no_mapped_name={skipped_no_mapped}"
    )
