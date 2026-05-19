import os
import openpyxl

from pages.models.epd import EPD, EPDImpact, Impact


EXCEL_PATH = os.path.join(
    os.path.dirname(__file__),
    "..", "..", "..", "docs",
    "2026-02-19 - Template_Generic_EPD_CORRECTED.xlsx",
)

TABS = ["India", "Indonesia", "Vietnam", "Cambodia", "Thailand"]


def update_generic_epd_gwp(stdout=None, style=None):
    def log(msg):
        if stdout:
            stdout.write(msg)

    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)

    gwp_impact = Impact.objects.filter(
        impact_category="gwp", life_cycle_stage="a1a3"
    ).first()

    if gwp_impact is None:
        raise RuntimeError("Impact with category='gwp' and stage='a1a3' not found in DB.")

    total_updated = 0
    total_skipped_not_found = 0
    total_skipped_no_uuid = 0

    for tab in TABS:
        if tab not in wb.sheetnames:
            log(f"  [WARN] Tab '{tab}' not found in workbook, skipping.")
            continue

        ws = wb[tab]
        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            log(f"  [WARN] Tab '{tab}' is empty, skipping.")
            continue

        headers = [str(h).strip() if h is not None else "" for h in rows[0]]

        try:
            uuid_col = headers.index("UUID")
        except ValueError:
            log(f"  [ERROR] Tab '{tab}' has no UUID column, skipping.")
            continue

        gwp_col = next(
            (i for i, h in enumerate(headers) if h.startswith("gwp_a1a3")),
            None,
        )
        if gwp_col is None:
            log(f"  [ERROR] Tab '{tab}' has no gwp_a1a3 column, skipping.")
            continue

        tab_updated = 0
        tab_skipped_not_found = 0
        tab_skipped_no_uuid = 0

        for row in rows[1:]:
            uuid_val = row[uuid_col]
            gwp_val = row[gwp_col]

            if uuid_val is None:
                tab_skipped_no_uuid += 1
                continue

            uuid_str = str(uuid_val).strip()
            if not uuid_str:
                tab_skipped_no_uuid += 1
                continue

            if gwp_val is None:
                tab_skipped_no_uuid += 1
                continue

            try:
                epd = EPD.objects.get(UUID=uuid_str)
            except EPD.DoesNotExist:
                tab_skipped_not_found += 1
                continue

            updated = EPDImpact.objects.filter(epd=epd, impact=gwp_impact).update(
                value=float(gwp_val)
            )

            if updated == 0:
                EPDImpact.objects.create(epd=epd, impact=gwp_impact, value=float(gwp_val))

            tab_updated += 1

        log(
            f"  [{tab}] updated={tab_updated}, "
            f"not_in_db={tab_skipped_not_found}, "
            f"skipped_null={tab_skipped_no_uuid}"
        )

        total_updated += tab_updated
        total_skipped_not_found += tab_skipped_not_found
        total_skipped_no_uuid += tab_skipped_no_uuid

    wb.close()

    log(
        f"Summary: updated={total_updated}, "
        f"not_in_db={total_skipped_not_found}, "
        f"skipped_null={total_skipped_no_uuid}"
    )
