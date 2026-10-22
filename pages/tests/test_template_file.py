"""The shipped template: valid for Excel, filled in, and importable.

One file, filled in as a worked example. A blank template plus a separate
example was tried and only raised the question of which to open.

The first test here is the one that matters most. Three different editors have
broken this workbook, and the worst break passed every check in the suite:

  openpyxl      destroys the 30 x14 extension validations - every dropdown.
  ElementTree   renames namespace prefixes. The worksheets carry
                mc:Ignorable="x14ac xr xr2 xr3"; ET rewrote it to
                ns1:Ignorable, renamed x14ac, and dropped xr, xr2 and xr3.
                The result is WELL-FORMED XML naming prefixes that no longer
                exist - so the zip opened, every part parsed, openpyxl read
                it, rows and cells were in order, and Excel still offered to
                repair it. Nine sheets shipped that way.
  a regex       `(?: [^>]*)?(?:/>|>.*?</c>)` lets [^>]* swallow the slash of
                a self-closing <c .../>, after which `.*?</c>` eats the next
                cell. Six building types vanished from an untouched column.
"""
import re
import warnings
import zipfile
from pathlib import Path

import openpyxl
import pytest

TEMPLATE = "static/docs/import-template.xlsx"
DIALOG = Path("templates/pages/home/import-dialog/step-upload.html")


def _book():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(TEMPLATE, data_only=True)


# --- the file must be valid for EXCEL, not merely well-formed -------------

def test_no_part_names_an_undeclared_namespace_prefix():
    """mc:Ignorable and the prefixes it lists must be declared on the same
    element. Naming a prefix that is not declared is what makes Excel offer
    to repair the workbook, and nothing else in this suite sees it."""
    archive = zipfile.ZipFile(TEMPLATE)
    problems = []
    for name in archive.namelist():
        if not name.endswith(".xml"):
            continue
        xml = archive.read(name).decode("utf-8", "ignore")
        body = xml[xml.index("?>") + 2:] if "?>" in xml else xml
        root = re.search(r"<\w+[^>]*?>", body)
        if root is None:
            continue
        root_tag = root.group(0)
        marker = re.search(r'([\w]+):Ignorable="([^"]+)"', root_tag)
        if marker:
            declared = set(re.findall(r"xmlns:([a-zA-Z0-9]+)=", root_tag))
            missing = [p for p in marker.group(2).split() if p not in declared]
            if missing:
                problems.append((name, "Ignorable names %s" % ",".join(missing)))
            if marker.group(1) not in declared:
                problems.append((name, "Ignorable prefix %s" % marker.group(1)))
        anywhere = set(re.findall(r"xmlns:([a-zA-Z0-9]+)=", xml))
        for prefix in set(re.findall(r"<(\w+):", xml)):
            if prefix != "xml" and prefix not in anywhere:
                problems.append((name, "element prefix %s:" % prefix))
    assert not problems, problems


def test_the_dropdowns_survived():
    archive = zipfile.ZipFile(TEMPLATE)
    dv = x14 = 0
    for name in archive.namelist():
        if not name.endswith(".xml"):
            continue
        xml = archive.read(name).decode("utf-8", "ignore")
        dv += len(re.findall(r"<dataValidation[ >]", xml))
        x14 += len(re.findall(r"<x14:dataValidation[ >]", xml))
    assert (len(archive.namelist()), dv, x14) == (48, 24, 30)


def test_rows_and_cells_are_in_order():
    """Excel rejects a sheet whose rows or cells are out of sequence."""
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    import xml.etree.ElementTree as ET

    def column(ref):
        n = 0
        for ch in "".join(c for c in ref if c.isalpha()):
            n = n * 26 + (ord(ch) - 64)
        return n

    archive = zipfile.ZipFile(TEMPLATE)
    for name in archive.namelist():
        if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
            continue
        data = ET.fromstring(archive.read(name)).find(ns + "sheetData")
        if data is None:
            continue
        last_row = 0
        for row in data.findall(ns + "row"):
            number = int(row.get("r"))
            assert number > last_row, (name, number)
            last_row = number
            last_col = 0
            for cell in row.findall(ns + "c"):
                col = column(cell.get("r"))
                assert col > last_col, (name, cell.get("r"))
                last_col = col


def test_the_workbook_opens():
    assert len(_book().sheetnames) == 26


# --- it is a worked example ----------------------------------------------

def test_it_is_filled_in():
    book = _book()
    filled = set()
    for sheet in book.sheetnames:
        if sheet == "List":
            continue
        ws = book[sheet]
        first = 3 if sheet.startswith("Structural") else 2
        for row in ws.iter_rows(min_row=first, values_only=True):
            values = [v for v in row if v not in (None, "")
                      and str(v).strip().lower() not in ("no", "yes")]
            if values:
                filled.add(sheet)
    for expected in ("Building Details", "Cooling System - Split air cond",
                     "Lightning System - LED", "Operational Energy Carriers",
                     "Structural Components - By Comp"):
        assert expected in filled, (expected, sorted(filled))


def test_neither_file_flag_demands_an_upload():
    """"Has certification" and "Has BoQ" each require a file when Yes. BoQ was
    Yes, so importing the example stopped at "At least one BoQ file is
    required" - a downloadable example has no files to attach."""
    details = _book()["Building Details"]
    assert str(details["I2"].value).strip().lower() == "no"
    assert str(details["J2"].value).strip().lower() == "no"


def test_the_structural_rows_are_whole():
    """Clearing this sheet once left components with no quantities and no
    materials, which reads as a broken file rather than an example."""
    ws = _book()["Structural Components - By Comp"]
    rows = [r for r in ws.iter_rows(min_row=3, values_only=True) if r[1]]
    assert len(rows) == 13
    for row in rows:
        assert row[3], row          # dimension
        assert row[4], row          # quantity
        assert row[8], row          # material


def test_the_hot_water_efficiency_is_a_percentage():
    """It was 3, which the importer divided by 0.03 - 33x the real energy."""
    assert float(_book()["Hot water System - Water Heater"]
                 .cell(row=2, column=9).value) >= 10


def test_the_fan_power_is_a_specific_fan_power():
    """75 W/CMH, read as a percentage, gave 4,322,970 kWh/yr from one row."""
    value = float(_book()["Ventillation System - Air Handl"]
                  .cell(row=2, column=2).value)
    assert 0.1 <= value <= 10


def test_the_efficiency_metric_is_stated():
    """Blank means COP. An ISEER machine stored as COP looks far more
    efficient than it is."""
    sheet = _book()["Cooling System - Split air cond"]
    assert str(sheet.cell(row=3, column=16).value).strip() == "ISEER"


def test_the_energy_intensity_is_believable():
    book = _book()
    area = float(book["Building Details"].cell(row=2, column=6).value)
    carriers = float(book["Operational Energy Carriers"]["C2"].value)
    assert 50 <= carriers / area <= 300, carriers / area


# --- the dialog ----------------------------------------------------------

def test_the_dialog_offers_one_download():
    html = DIALOG.read_text(encoding="utf-8")
    assert "import_template_url" in html
    assert "import_example_url" not in html
    assert "Download sample template" in html


def test_the_dialog_says_to_type_over_the_example():
    """A filled file with no explanation invites importing someone else's
    building by accident."""
    html = DIALOG.read_text(encoding="utf-8")
    assert "Replace each value with your own" in html
    assert "imported as yours" in html


def test_there_is_no_second_file():
    assert not Path("static/docs/import-template-example.xlsx").exists()
