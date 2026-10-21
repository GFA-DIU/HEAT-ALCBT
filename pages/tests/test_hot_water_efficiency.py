"""A water heater that imported 33x its real energy, from the template's own
example value.

"Baseline hot water system equipment efficiency level" is a PERCENTAGE - the
manual form asks for 75-95 and caps it at 100 - but the header said only
"efficiency level", and the shipped template's example row carried 3.
"""
import warnings

import openpyxl
import pytest

from pages.models.building_operation import HotWaterSystemType
from pages.views.building.import_building import _calc_hws_energy

TEMPLATE = "static/docs/import-template.xlsx"
# The blank template carries no values now - it is a shape to fill. The worked
# example is where a figure lives, and it is the one people copy.
EXAMPLE = "static/docs/import-template-example.xlsx"

# 6 units x 3.5 kW x 2 h/day x 6 days x 50 weeks = 12,600 kWh before efficiency
ARGS = (3.5, 6, 2, 6, 50)


def test_a_sensible_percentage_gives_a_sensible_figure():
    energy = _calc_hws_energy(HotWaterSystemType.SOLAR, *ARGS, 95, "electricity")
    assert 13000 < energy < 13500


def test_the_value_the_template_used_to_teach_is_refused():
    """3 was the example. 12,600 / 0.03 = 420,000 kWh/yr - a plausible number
    for a large building, and 33x the truth."""
    assert _calc_hws_energy(HotWaterSystemType.SOLAR, *ARGS, 3, "electricity") is None


def test_a_fraction_is_refused_too():
    """0.85 passes the form's own min of 0.01 and would divide by 0.0085."""
    assert _calc_hws_energy(HotWaterSystemType.SOLAR, *ARGS, 0.85, "electricity") is None


@pytest.mark.parametrize("eff", [10, 75, 95, 100])
def test_plausible_efficiencies_still_compute(eff):
    assert _calc_hws_energy(HotWaterSystemType.SOLAR, *ARGS, eff, "electricity")


def test_the_heat_pump_path_is_untouched():
    """Heat pumps divide by COP, a dimensionless 2.5-4.5, not a percentage.
    The guard must not reject a COP of 3."""
    energy = _calc_hws_energy(HotWaterSystemType.HEAT_PUMP, *ARGS, 3.0)
    assert energy == 4200


def test_a_solar_system_on_solar_fuel_is_still_zero():
    assert _calc_hws_energy(HotWaterSystemType.SOLAR, *ARGS, 95, "solar") == 0


def _sheet(name, path=TEMPLATE):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path, data_only=True)[name]


def test_the_template_says_the_column_is_a_percentage():
    ws = _sheet("Hot water System - Water Heater")
    header = str(ws.cell(row=1, column=9).value)
    assert "%" in header, header


def test_the_worked_examples_efficiency_is_plausible():
    """It was 3 in the template that shipped. Anyone copying that row got a
    water heater reported at 33x its energy."""
    ws = _sheet("Hot water System - Water Heater", EXAMPLE)
    assert float(ws.cell(row=2, column=9).value) >= 10


def test_the_blank_template_carries_no_value_to_copy():
    """A blank template cannot teach a wrong number."""
    ws = _sheet("Hot water System - Water Heater")
    assert ws.cell(row=2, column=9).value in (None, "")
