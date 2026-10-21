"""The template must offer names the catalogue still uses.

Migration 0046 consolidated the taxonomy. The template went on offering the
old names for months: 42 of its 90 construction techniques and 15 of its 19
building components no longer existed. They imported through the legacy maps,
so nothing failed - a user was simply picking from a list of retired names.

The maps stay, because files saved from the old template still carry the old
wording. This is about what a NEW file offers.
"""
import re
import warnings
import zipfile
from pathlib import Path

import openpyxl
import pytest

TEMPLATE = "static/docs/import-template.xlsx"


@pytest.fixture(scope="module")
def workbook():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(TEMPLATE, data_only=True)


@pytest.fixture(scope="module")
def dropdown_ranges():
    """Every List range a dropdown points at, read from the x14 validations."""
    archive = zipfile.ZipFile(TEMPLATE)
    found = set()
    for name in archive.namelist():
        if not name.startswith("xl/worksheets/"):
            continue
        xml = archive.read(name).decode("utf-8", "ignore")
        found.update(re.findall(
            r"<xm:f>List!\$([A-Z]+)\$(\d+):\$[A-Z]+\$(\d+)</xm:f>", xml))
    return sorted(found)


def _values(sheet, column, first, last):
    out = []
    for row in range(int(first), int(last) + 1):
        value = sheet["%s%d" % (column, row)].value
        if value not in (None, ""):
            out.append(str(value).strip())
    return out


# --- the file itself -------------------------------------------------------

def test_the_validations_survived_the_edit():
    """openpyxl cannot round-trip this workbook - saving through it destroys
    the x14 extension validations, which is every dropdown. The lists were
    rewritten as XML in place for that reason."""
    archive = zipfile.ZipFile(TEMPLATE)
    counts = {"dataValidation": 0, "x14:dataValidation": 0}
    for name in archive.namelist():
        if not name.endswith(".xml"):
            continue
        xml = archive.read(name).decode("utf-8", "ignore")
        counts["dataValidation"] += len(re.findall(r"<dataValidation[ >]", xml))
        counts["x14:dataValidation"] += len(
            re.findall(r"<x14:dataValidation[ >]", xml))
    assert len(archive.namelist()) == 48
    assert counts["dataValidation"] == 24
    assert counts["x14:dataValidation"] == 30


def test_the_workbook_still_opens(workbook):
    assert len(workbook.sheetnames) == 26


# --- what the dropdowns offer ---------------------------------------------
#
# The consolidated taxonomy is not created by a migration, so "is the table
# empty" would let these pass on a database holding the PRE-0046 names - which
# is how a test ends up unable to fail. This name exists only after the
# consolidation.
_SENTINEL = "Plain Cement Concrete (PCC) Footings"


def _require_consolidated_taxonomy():
    from pages.models.assembly import AssemblyTechnique

    if not AssemblyTechnique.objects.filter(name__iexact=_SENTINEL).exists():
        pytest.skip(
            "this database does not hold the consolidated (post-0046) "
            "taxonomy - %r is absent" % _SENTINEL)


@pytest.mark.django_db
def test_every_technique_offered_exists(workbook, dropdown_ranges):
    from pages.models.assembly import AssemblyTechnique

    _require_consolidated_taxonomy()
    known = {t.lower() for t in
             AssemblyTechnique.objects.values_list("name", flat=True)}

    sheet = workbook["List"]
    stale = set()
    for column, first, last in dropdown_ranges:
        if column != "W":
            continue
        for value in _values(sheet, column, first, last):
            if value.lower() not in known:
                stale.add(value)
    assert not stale, "techniques offered that no longer exist: %s" % sorted(stale)


@pytest.mark.django_db
def test_every_component_offered_exists(workbook):
    from pages.views.building.import_building import _get_structural_category_map

    _require_consolidated_taxonomy()
    cat_map = _get_structural_category_map()

    sheet = workbook["List"]
    offered = _values(sheet, "I", 2, 15)
    assert len(offered) == 14, offered
    stale = [c for c in offered if c.lower() not in cat_map]
    assert not stale, stale


@pytest.mark.django_db
def test_the_components_are_the_consolidated_fourteen(workbook):
    from pages.models.assembly import AssemblyCategory

    _require_consolidated_taxonomy()
    current = {c.name for c in AssemblyCategory.objects.all()}

    offered = set(_values(workbook["List"], "I", 2, 15))
    assert offered == current, {
        "offered not in catalogue": sorted(offered - current),
        "in catalogue not offered": sorted(current - offered),
    }


# --- the example rows must agree with the lists ----------------------------

@pytest.mark.django_db
def test_the_example_rows_use_names_the_dropdowns_offer(workbook):
    """Updating the lists without the examples would ship a file whose own
    sample data contradicts its pick-lists."""
    from pages.models.assembly import AssemblyTechnique

    _require_consolidated_taxonomy()
    known = {t.lower() for t in
             AssemblyTechnique.objects.values_list("name", flat=True)}

    sheet = workbook["Structural Components - By Comp"]
    components = set(_values(workbook["List"], "I", 2, 15))

    problems = []
    for row in range(3, sheet.max_row + 1):
        component = sheet.cell(row=row, column=2).value
        if not component:
            continue
        component = str(component).strip()
        technique = str(sheet.cell(row=row, column=3).value or "").strip()
        if component not in components:
            problems.append((row, "component", component))
        if technique and technique.lower() not in known:
            problems.append((row, "technique", technique))
    assert not problems, problems


# --- the legacy wording must still import ---------------------------------

@pytest.mark.django_db
def test_the_old_names_are_still_accepted():
    """Removing them from the list must not stop files already saved with
    them - including the Cambodia 250-building set."""
    from pages.views.building.import_building import (
        _get_structural_category_map, _resolve_technique_name,
    )

    _require_consolidated_taxonomy()
    cat_map = _get_structural_category_map()

    assert cat_map["bottom floor construction"].name == "Ground Floor"
    assert cat_map["exterior walls"].name == "External Walls"
    finishes = cat_map["floor finish"]
    assert _resolve_technique_name(finishes, "Stones and Tiles") == "Stone & Tiles"
