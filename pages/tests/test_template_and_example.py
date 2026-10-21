"""Two downloads: a blank template, and a worked example.

The template used to ship with 23 rows of example data, so one file was both
the thing you fill in and the example you learn from. Anything a user forgot
to delete imported as their own building - and the template's own example
named a German reinforcement EPD, a hot-water efficiency of 3 where the column
is a percentage, and components migration 0046 had removed.
"""
import re
import warnings
import zipfile
from pathlib import Path

import openpyxl
import pytest

BLANK = "static/docs/import-template.xlsx"
EXAMPLE = "static/docs/import-template-example.xlsx"
DIALOG = Path("templates/pages/home/import-dialog/step-upload.html")


def _load(path):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path, data_only=True)


def _census(path):
    archive = zipfile.ZipFile(path)
    dv = x14 = 0
    for name in archive.namelist():
        if not name.endswith(".xml"):
            continue
        xml = archive.read(name).decode("utf-8", "ignore")
        dv += len(re.findall(r"<dataValidation[ >]", xml))
        x14 += len(re.findall(r"<x14:dataValidation[ >]", xml))
    return len(archive.namelist()), dv, x14


def _filled_rows(workbook):
    """Rows holding anything beyond the pre-filled yes/no defaults."""
    out = []
    for sheet in workbook.sheetnames:
        if sheet == "List":
            continue
        ws = workbook[sheet]
        first = 3 if sheet.startswith("Structural") else 2
        for row in ws.iter_rows(min_row=first, values_only=True):
            values = [v for v in row if v not in (None, "")
                      and str(v).strip().lower() not in ("no", "yes")]
            if values:
                out.append((sheet, values[:3]))
    return out


# --- both files keep their dropdowns --------------------------------------

@pytest.mark.parametrize("path", [BLANK, EXAMPLE])
def test_the_validations_survived(path):
    """openpyxl cannot save this workbook - it destroys the x14 extension
    validations, which is every dropdown. Both files were edited as XML."""
    assert _census(path) == (48, 24, 30)


@pytest.mark.parametrize("path", [BLANK, EXAMPLE])
def test_the_workbook_opens(path):
    assert len(_load(path).sheetnames) == 26


def test_the_example_exists_as_its_own_file():
    assert Path(EXAMPLE).exists()


# --- the blank template is blank ------------------------------------------

def test_the_blank_template_has_no_example_data():
    """23 rows used to ship in it. A user who missed one imported someone
    else's building alongside their own."""
    left = _filled_rows(_load(BLANK))
    assert not left, "example data left in the blank template: %s" % left[:5]


def test_the_blank_template_keeps_its_headers():
    wb = _load(BLANK)
    assert wb["Lightning System - LED"]["A1"].value.strip() == "Room type"
    assert "Efficiency metric" in str(
        wb["Cooling System - Split air cond"].cell(row=1, column=16).value)


def test_the_blank_template_keeps_its_dropdown_lists():
    """Clearing the example rows must not touch the List sheet."""
    sheet = _load(BLANK)["List"]
    types = [sheet["B%d" % r].value for r in range(2, 53)]
    assert len([t for t in types if t]) == 42
    components = [sheet["I%d" % r].value for r in range(2, 16)]
    assert len([c for c in components if c]) == 14


# --- the worked example is complete and correct ---------------------------

def test_the_example_is_filled():
    filled = _filled_rows(_load(EXAMPLE))
    sheets = {s for s, _ in filled}
    for expected in ("Building Details", "Cooling System - Split air cond",
                     "Lightning System - LED", "Operational Energy Carriers",
                     "Structural Components - By Comp"):
        assert expected in sheets, (expected, sorted(sheets))


def test_the_example_hot_water_efficiency_is_a_percentage():
    """The old example said 3, which the importer divided by 0.03 - a water
    heater reported at 33x its energy."""
    value = _load(EXAMPLE)["Hot water System - Water Heater"].cell(
        row=2, column=9).value
    assert float(value) >= 10


def test_the_example_fan_power_is_a_specific_fan_power():
    """75 W/CMH, read as a percentage, gave 4,322,970 kWh/yr from one row."""
    value = _load(EXAMPLE)["Ventillation System - Air Handl"].cell(
        row=2, column=2).value
    assert 0.1 <= float(value) <= 10


def test_the_example_states_its_efficiency_metric():
    """Blank means COP. An ISEER machine stored as COP looks far more
    efficient than it is."""
    sheet = _load(EXAMPLE)["Cooling System - Split air cond"]
    assert str(sheet.cell(row=3, column=16).value).strip() == "ISEER"


@pytest.mark.django_db
def test_the_example_building_type_resolves():
    from pages.views.building.import_building import _resolve_building_type

    value = _load(EXAMPLE)["Building Details"]["B2"].value
    resolved, error = _resolve_building_type(str(value).strip())
    if resolved is None and "not found" in str(error):
        pytest.skip("building categories are not loaded in this database")
    assert resolved is not None, error


@pytest.mark.django_db
def test_the_example_carriers_account_for_its_systems():
    """The import refuses carriers more than 5% from the systems. The example
    must not be a file that its own importer rejects.

    Needs the database because parsing a lighting row resolves its room type
    against LightingSpaceType. The energy figures do not depend on that, so
    the arithmetic below holds whether or not the table is populated.
    """
    from pages.views.building import import_building as ib
    from pages.models.building_operation.ventilation import VentilationType
    from pages.models.building_operation import HotWaterSystemType

    wb = _load(EXAMPLE)

    def rows(sheet, vent=False):
        out = []
        for row in wb[sheet].iter_rows(min_row=2, values_only=True):
            row = [None if v is None else str(v).strip() for v in row]
            keep = (row and row[0] not in (None, "")) if vent else any(
                c not in (None, "") and str(c).lower() not in
                ("none", "no", "yes") for c in row)
            if keep:
                out.append(row)
        return out

    total = 0.0
    for parsed, key in (
        (ib._parse_cooling_rows("split ac", rows("Cooling System - Split air cond")),
         "total_energy_consumption_annually"),
        (ib._parse_cooling_rows("vrf", rows("Cooling System - Variable refri")),
         "total_energy_consumption_annually"),
        (ib._parse_ventilation_rows(VentilationType.AHU,
                                    rows("Ventillation System - Air Handl", True)),
         "annual_energy_consumption"),
        (ib._parse_ventilation_rows(VentilationType.CASSETTE_AC,
                                    rows("Ventillation System - CeilingWa", True)),
         "annual_energy_consumption"),
        (ib._parse_lighting_rows("led", rows("Lightning System - LED")),
         "total_energy_consumption_kwh_per_year"),
        (ib._parse_hot_water_rows(HotWaterSystemType.SOLAR,
                                  rows("Hot water System - Water Heater")),
         "total_energy_consumption_kwh_per_year"),
    ):
        for row in parsed:
            total += float(row[key])

    carriers = float(wb["Operational Energy Carriers"]["C2"].value)
    assert total > 0
    assert abs(carriers - total) / total <= 0.05, (carriers, total)


def test_the_example_energy_intensity_is_believable():
    """2,541 kWh/m2/yr was the figure before the ventilation guard. A Phnom
    Penh office is nearer 100-250."""
    wb = _load(EXAMPLE)
    area = float(wb["Building Details"].cell(row=2, column=6).value)
    carriers = float(wb["Operational Energy Carriers"]["C2"].value)
    assert 50 <= carriers / area <= 300, carriers / area


# --- the dialog offers both -----------------------------------------------

def test_the_dialog_offers_both_downloads():
    html = DIALOG.read_text(encoding="utf-8")
    assert "import_template_url" in html
    assert "import_example_url" in html
    assert "Download blank template" in html
    assert "See a worked example" in html


def test_the_dialog_says_which_is_which():
    """Without a word of explanation, two links are worse than one."""
    html = DIALOG.read_text(encoding="utf-8")
    assert "not meant to be uploaded as-is" in html


def test_the_view_supplies_both_urls():
    import inspect

    from pages.views.building import import_building as ib

    src = inspect.getsource(ib.view_import_building_step)
    assert "docs/import-template.xlsx" in src
    assert "docs/import-template-example.xlsx" in src
